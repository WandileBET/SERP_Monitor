from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from pathlib import Path

from openpyxl import load_workbook

from fastapi import Depends, FastAPI, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy import desc, func, select, text
from sqlalchemy.orm import Session

from .config import get_settings
from .db import engine, get_db
from .models import Capture, Keyword, SerpResult
from .schemas import CaptureSummary, DashboardOut, KeywordOut, LatestResult, TrendPoint

settings = get_settings()
app = FastAPI(title=settings.app_name)
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.cors_origin_list,
    allow_credentials=False,
    allow_methods=["GET"],
    allow_headers=["Content-Type"],
)

LOCAL_TZ = ZoneInfo("Africa/Johannesburg")
TARGET_DOMAIN = "hollywoodbets.net"


def normalize_keyword(value: str) -> str:
    return " ".join(value.split()).strip()


def get_keyword(db: Session, keyword: str) -> Keyword:
    normalized = normalize_keyword(keyword)
    row = db.scalar(select(Keyword).where(Keyword.keyword == normalized))
    if not row:
        raise HTTPException(status_code=404, detail="Keyword not found")
    return row


def is_hollywoodbets(domain: str | None, name: str | None = None) -> bool:
    values = [(domain or "").lower().strip().removeprefix("www."), (name or "").lower().strip()]
    return any(value == TARGET_DOMAIN or value.endswith("." + TARGET_DOMAIN) or "hollywoodbets" in value for value in values)


def parse_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"Invalid date: {value}") from exc


def local_bounds(start_date: date | None, end_date: date | None, latest_at: datetime | None):
    if latest_at is None and (start_date is None or end_date is None):
        return None, None, None, None

    latest_local = latest_at.astimezone(LOCAL_TZ) if latest_at else None
    resolved_end = end_date or latest_local.date()
    resolved_start = start_date or resolved_end
    if resolved_start > resolved_end:
        raise HTTPException(status_code=400, detail="Start date cannot be after end date.")

    start_local = datetime.combine(resolved_start, datetime.min.time(), tzinfo=LOCAL_TZ)
    end_local = datetime.combine(resolved_end + timedelta(days=1), datetime.min.time(), tzinfo=LOCAL_TZ)
    return start_local, end_local, resolved_start.isoformat(), resolved_end.isoformat()


def choose_resolution(start_iso: str, end_iso: str, requested: str) -> str:
    span_days = (date.fromisoformat(end_iso) - date.fromisoformat(start_iso)).days + 1
    if requested != "auto":
        return requested
    if span_days <= 2:
        return "hourly"
    if span_days <= 45:
        return "daily"
    if span_days <= 180:
        return "weekly"
    return "monthly"


def bucket_start(captured_at: datetime, resolution: str) -> datetime:
    local = captured_at.astimezone(LOCAL_TZ)
    if resolution == "daily":
        return datetime.combine(local.date(), datetime.min.time(), tzinfo=LOCAL_TZ)
    if resolution == "weekly":
        day = local.date() - timedelta(days=local.weekday())
        return datetime.combine(day, datetime.min.time(), tzinfo=LOCAL_TZ)
    if resolution == "monthly":
        return datetime(local.year, local.month, 1, tzinfo=LOCAL_TZ)
    return captured_at


def build_trend(rows: list[tuple[SerpResult, datetime]], resolution: str) -> list[TrendPoint]:
    if resolution == "hourly":
        return [
            TrendPoint(captured_at=dt, rank=result.rank, name=result.name, domain=result.domain)
            for result, dt in rows
        ]

    # For aggregated ranges, use the latest observed rank in each bucket for each domain/name.
    buckets: dict[tuple[datetime, str], tuple[datetime, SerpResult]] = {}
    for result, captured_at in rows:
        identity = (result.domain or result.name or "").lower()
        if not identity:
            continue
        bucket = bucket_start(captured_at, resolution)
        key = (bucket, identity)
        previous = buckets.get(key)
        if previous is None or captured_at > previous[0]:
            buckets[key] = (captured_at, result)

    output: list[TrendPoint] = []
    for (bucket, _identity), (captured_at, result) in sorted(buckets.items(), key=lambda item: (item[0][0], item[0][1])):
        output.append(
            TrendPoint(
                captured_at=bucket,
                rank=result.rank,
                name=result.name,
                domain=result.domain,
            )
        )
    return output


def build_heatmap(rows: list[tuple[SerpResult, datetime]], requested_date: date | None = None) -> tuple[list[TrendPoint], str | None]:
    if not rows:
        return [], requested_date.isoformat() if requested_date else None
    available_dates = {dt.astimezone(LOCAL_TZ).date() for _, dt in rows}
    latest_local_date = requested_date if requested_date in available_dates else max(available_dates)
    latest_day_rows = [
        (result, dt) for result, dt in rows if dt.astimezone(LOCAL_TZ).date() == latest_local_date
    ]
    # Keep the best rank for a domain at each hour when duplicate domain pages appear.
    best: dict[tuple[int, str], TrendPoint] = {}
    for result, dt in latest_day_rows:
        hour = dt.astimezone(LOCAL_TZ).hour
        identity = (result.domain or result.name or "").lower()
        key = (hour, identity)
        current = best.get(key)
        point = TrendPoint(captured_at=dt, rank=result.rank, name=result.name, domain=result.domain)
        if current is None or result.rank < current.rank:
            best[key] = point
    return list(best.values()), latest_local_date.isoformat()


@app.on_event("startup")
def startup() -> None:
    with engine.connect() as conn:
        conn.execute(text("SELECT 1"))


@app.get("/health")
def health(db: Session = Depends(get_db)):
    db.execute(text("SELECT 1"))
    return {"status": "ok"}


@app.get("/api/keywords", response_model=list[KeywordOut])
def keywords(db: Session = Depends(get_db)):
    return list(
        db.scalars(
            select(Keyword)
            .where(Keyword.active.is_(True))
            .order_by(Keyword.keyword.asc())
        )
    )


@app.get("/api/dashboard", response_model=DashboardOut)
def dashboard(
    keyword: str = Query(min_length=1),
    start_date: str | None = Query(default=None),
    end_date: str | None = Query(default=None),
    resolution: str = Query(default="hourly", pattern="^(auto|hourly|daily|weekly|monthly)$"),
    hour: int | None = Query(default=None, ge=0, le=23),
    heatmap_date: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    kw = get_keyword(db, keyword)
    start_requested = parse_date(start_date)
    end_requested = parse_date(end_date)
    heatmap_requested = parse_date(heatmap_date)

    first_capture = db.scalar(
        select(Capture.captured_at)
        .where(Capture.keyword_id == kw.id)
        .order_by(Capture.captured_at.asc())
        .limit(1)
    )
    latest_overall = db.scalar(
        select(Capture.captured_at)
        .where(Capture.keyword_id == kw.id)
        .order_by(Capture.captured_at.desc())
        .limit(1)
    )

    start_local, end_local, resolved_start, resolved_end = local_bounds(
        start_requested, end_requested, latest_overall
    )
    if start_local is None or end_local is None:
        return DashboardOut(
            keyword=kw.keyword,
            report_start=None,
            report_end=None,
            report_date=None,
            resolution="hourly",
            hour_filter=hour,
            capture_count=0,
            daily_capture_count=0,
            first_capture_at=first_capture,
            latest_capture=None,
            latest_results=[],
            trend=[],
            heatmap=[],
            heatmap_date=None,
        )

    effective_resolution = choose_resolution(resolved_start, resolved_end, resolution)

    captures_query = (
        select(Capture)
        .where(
            Capture.keyword_id == kw.id,
            Capture.captured_at >= start_local,
            Capture.captured_at < end_local,
        )
        .order_by(Capture.captured_at.asc())
    )
    captures_in_range = list(db.scalars(captures_query))
    if hour is not None:
        captures_in_range = [
            row for row in captures_in_range
            if row.captured_at.astimezone(LOCAL_TZ).hour == hour
        ]

    latest = captures_in_range[-1] if captures_in_range else None
    latest_results: list[LatestResult] = []
    latest_summary = None
    latest_rows: list[SerpResult] = []

    if latest:
        latest_rows = list(
            db.scalars(
                select(SerpResult)
                .where(SerpResult.capture_id == latest.id)
                .order_by(SerpResult.rank.asc())
            )
        )
        latest_results = [
            LatestResult(rank=r.rank, name=r.name, domain=r.domain, url=r.url, title=r.title)
            for r in latest_rows
        ]
        latest_summary = CaptureSummary(
            id=latest.id,
            keyword=kw.keyword,
            captured_at=latest.captured_at,
            result_count=len(latest_rows),
        )

    capture_ids = [row.id for row in captures_in_range]
    raw_rows: list[tuple[SerpResult, datetime]] = []
    if capture_ids:
        result_rows = db.execute(
            select(SerpResult, Capture.captured_at)
            .join(Capture, Capture.id == SerpResult.capture_id)
            .where(SerpResult.capture_id.in_(capture_ids))
            .order_by(Capture.captured_at.asc(), SerpResult.rank.asc())
        )
        raw_rows = list(result_rows)

    trend = build_trend(raw_rows, effective_resolution)
    heatmap, selected_heatmap_date = build_heatmap(raw_rows, heatmap_requested)

    hb_ranks = [
        result.rank for result, _dt in raw_rows if is_hollywoodbets(result.domain, result.name)
    ]
    top3_capture_count = sum(
        1 for capture_id in capture_ids
        if any(
            result.rank <= 3 and is_hollywoodbets(result.domain, result.name)
            for result, _dt in raw_rows
            if result.capture_id == capture_id
        )
    )

    competitors = {
        (row.domain or row.name).lower()
        for row in latest_rows
        if not is_hollywoodbets(row.domain, row.name)
    }

    return DashboardOut(
        keyword=kw.keyword,
        report_start=resolved_start,
        report_end=resolved_end,
        report_date=latest.captured_at.astimezone(LOCAL_TZ).date().isoformat() if latest else None,
        resolution=effective_resolution,
        hour_filter=hour,
        capture_count=len(captures_in_range),
        daily_capture_count=len(captures_in_range),
        first_capture_at=first_capture,
        latest_capture=latest_summary,
        latest_results=latest_results,
        trend=trend,
        heatmap=heatmap,
        heatmap_date=selected_heatmap_date,
        hollywoodbets_current_rank=next(
            (row.rank for row in latest_rows if is_hollywoodbets(row.domain, row.name)), None
        ),
        hollywoodbets_best_rank=min(hb_ranks) if hb_ranks else None,
        hollywoodbets_worst_rank=max(hb_ranks) if hb_ranks else None,
        hollywoodbets_average_rank=round(sum(hb_ranks) / len(hb_ranks), 1) if hb_ranks else None,
        competitors_tracked=len(competitors),
        top3_capture_count=top3_capture_count,
    )



@app.get("/api/monitor-status")
def monitor_status():
    """Return the latest capture status for each keyword workbook."""
    project_root = Path(__file__).resolve().parents[2]
    scripts_dir = Path(settings.serp_scripts_dir or (project_root / "scripts"))
    statuses = []

    if not scripts_dir.exists():
        return {"items": [], "bot_detected": []}

    for workbook_path in sorted(scripts_dir.glob("*_rankings.xlsx")):
        keyword = workbook_path.stem[:-len("_rankings")].replace("_", " ").strip()
        if not workbook_path.is_file():
            continue
        try:
            wb = load_workbook(workbook_path, read_only=True, data_only=True)
            if "Capture Status" not in wb.sheetnames:
                wb.close()
                continue
            ws = wb["Capture Status"]
            latest = None
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 7 or not row[0]:
                    continue
                date_value, hour_value, attempt, timestamp, status, valid_results, details = row[:7]
                candidate = timestamp or date_value
                if latest is None or str(candidate) > str(latest[0]):
                    latest = (candidate, date_value, hour_value, attempt, status, valid_results, details)
            wb.close()
            if latest is None:
                continue
            candidate, date_value, hour_value, attempt, status, valid_results, details = latest
            statuses.append({
                "keyword": keyword,
                "status": str(status or ""),
                "date": date_value.isoformat() if hasattr(date_value, "isoformat") else str(date_value or ""),
                "hour": str(hour_value or ""),
                "attempt": int(attempt) if isinstance(attempt, (int, float)) else str(attempt or ""),
                "valid_results": int(valid_results) if isinstance(valid_results, (int, float)) else 0,
                "details": str(details or ""),
            })
        except Exception:
            continue

    bots = [item for item in statuses if item["status"].strip().lower() == "bot detected"]
    return {"items": statuses, "bot_detected": bots}

@app.get("/api/captures", response_model=list[CaptureSummary])
def captures(
    keyword: str = Query(min_length=1),
    limit: int = Query(default=168, ge=1, le=2000),
    db: Session = Depends(get_db),
):
    kw = get_keyword(db, keyword)
    rows = list(
        db.scalars(
            select(Capture)
            .where(Capture.keyword_id == kw.id)
            .order_by(desc(Capture.captured_at))
            .limit(limit)
        )
    )
    return [
        CaptureSummary(
            id=row.id,
            keyword=kw.keyword,
            captured_at=row.captured_at,
            result_count=len(row.results),
        )
        for row in rows
    ]
